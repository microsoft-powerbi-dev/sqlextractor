"""Mapping Studio: a localhost-only JSON API over control.db plus a static drag-and-drop page.

Every /api request needs the per-launch token (X-Studio-Token). The Host header must be loopback,
which blocks DNS-rebinding pages from reaching the API. Errors carry only ExportError text.
"""
import hashlib
import hmac
import io
import json
import re
import secrets
import threading
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from .. import __version__, catalog
from ..common import ExportError, canonical, require
from ..config import check_config, load_config
from ..control import Control
from ..engine import export, make_plan
from ..mapping import HEADERS, load_mapping, normalize_model
from ..mapping_jobs import start_job, status_job, cancel_job
from ..sources import open_source, select_sql
from ..targets import automap, import_target
from ..transforms import FIELDS, TYPES, compile_columns

STATIC = Path(__file__).resolve().parent / "static"
ASSETS = {"/": ("index.html", "text/html; charset=utf-8"), "/app.js": ("app.js", "text/javascript; charset=utf-8"),
          "/app.css": ("app.css", "text/css; charset=utf-8")}
LIMIT = 10 * 1024 * 1024


def workbook(model):
    import openpyxl
    book = openpyxl.Workbook()
    columns = book.active
    columns.title = "Columns"
    columns.append(HEADERS)
    for column in model["columns"]:
        columns.append([canonical(column[k]) if k == "rules" else column[k] for k in HEADERS])
    meta = book.create_sheet("Metadata")
    meta.append(["key", "value"])
    for key, value in model["metadata"].items():
        meta.append([key, value])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class Studio:
    def __init__(self, workspace):
        self.control = Control(workspace)
        self.root = self.control.root
        self.lock = threading.Lock()
        self.running = set()

    def state(self, _):
        return {"version": __version__, "workspace": str(self.root), "connections": self.control.connections(),
                "sources": self.control.sources(), "targets": self.control.targets(), "mappings": self.control.mappings(),
                "contract": {"types": sorted(TYPES), "operations": {k: sorted(v) for k, v in FIELDS.items()}}}

    def save_connection(self, body):
        return {"id": self.control.save_connection(body.get("name"), body.get("connection_env"))}

    def _env(self, connection_id):
        return self.control.connection(connection_id)["connection_env"]

    def tables(self, body):
        return catalog.list_tables(self._env(body.get("connection_id")), body.get("schema") or None)

    def describe(self, body):
        kind = body.get("kind")
        if kind == "jsonl":
            return catalog.describe_fixture(self._fixture(body.get("fixture")))
        if kind == "query":
            return catalog.describe_query(self._env(body.get("connection_id")), body.get("query") or "")
        return catalog.describe_table(self._env(body.get("connection_id")), body.get("schema"), body.get("table"))

    def _fixture(self, path):
        require(isinstance(path, str) and path.strip(), "Fixture path is required")
        path = Path(path).expanduser().resolve()
        require(path.is_file() and path.suffix.lower() == ".jsonl", "Fixture must be an existing .jsonl file")
        return path

    def save_source(self, body):
        kind = body.get("kind")
        fixture = str(self._fixture(body.get("fixture"))) if kind == "jsonl" else None
        metadata = body.get("metadata")
        if metadata is None and body.get("capture", True):
            metadata = self.describe(body | {"fixture": fixture})
        pk = self.control.save_source(body.get("name"), kind, body.get("date_column"), body.get("connection_id"),
                                      body.get("schema") or None, body.get("table") or None, body.get("query") or None,
                                      fixture, body.get("settings") or {}, metadata)
        return self.control.source(pk)

    def save_target(self, body):
        return self.control.target(self.control.save_target(body.get("spec"), body.get("origin"), body.get("origin_sha256")))

    def upload(self, kind, name, data):
        require(kind in {"target", "mapping", "document"}, "Unknown upload kind")
        require(0 < len(data) <= LIMIT, "Uploads must be 1 byte to 10 MiB")
        stem = re.sub(r"[^A-Za-z0-9._-]", "_", Path(name or "upload").name)[-80:].lstrip(".") or "upload"
        path = self.root / "uploads" / (hashlib.sha256(data).hexdigest()[:12] + "-" + stem)
        if not path.exists():
            path.write_bytes(data)
        if kind == "target":
            result = import_target(path)
            pk = self.control.save_target(result["spec"], result["origin"], result["origin_sha256"])
            return {"target": self.control.target(pk), "warnings": result["warnings"]}
        if kind == "mapping":
            require(path.suffix.lower() in {".xlsx", ".xls", ".json"}, "Contract mappings are .xlsx, .xls, or .json; send business documents to Devin instead")
            return {"model": load_mapping(path)["model"]}
        return {"document": path.name}

    def automap(self, body):
        source = self.control.source(body.get("source_id"))
        spec = self.control.target(body.get("target_id"))["spec"]
        return automap(spec, [c["name"] for c in source["metadata"] or []], body.get("mapping_id") or None)

    def validate(self, body):
        model = normalize_model(body.get("model"))
        return {"model": model, "issues": self.control.review(model, body.get("source_id"), body.get("target_id"))}

    def save_mapping(self, body):
        return self.control.save_mapping(body.get("model"), body.get("source_id"), body.get("target_id"), "human", body.get("id"))

    def mapping(self, pk):
        return self.control.mapping(pk)

    def approve(self, pk, _):
        return self.control.approve(pk, "human")

    def retire(self, pk, _):
        self.control.retire(pk, "human")
        return {"id": pk, "status": "retired"}

    def _config(self, row):
        _, config = self.control.build_config(row)
        return check_config(config, self.root / "configs")

    def dry_run(self, pk, _):
        row = self.control.mapping(pk)
        require(row["source_id"] is not None and row["target_id"] is not None, "Link a source and a target first")
        config = self._config(row)
        plan = make_plan(config, {"model": row["model"]})
        return {"sql": select_sql(config, {"model": row["model"]}), "windows": len(plan["windows"]),
                "files": [w["name"] + config["file_extension"] for w in plan["windows"][:5]], "start": plan["start"], "end": plan["end"]}

    def preview(self, body):
        """Offline fixtures only: SQL rows never reach the browser."""
        model = normalize_model(body.get("model"))
        source = self.control.source(body.get("source_id"))
        require(source["kind"] == "jsonl", "Preview runs on offline JSONL fixture sources only")
        compiled = compile_columns(model["columns"])
        rows, errors = [], []
        with Path(source["fixture"]).open(encoding="utf-8") as stream:
            for number, line in enumerate(stream, 1):
                if len(rows) + len(errors) >= 20:
                    break
                if not line.strip():
                    continue
                try:
                    rows.append([run(json.loads(line)) for run in compiled])
                except ExportError as exc:
                    errors.append(f"Fixture row {number}: {exc}")
        return {"header": [c["target"] for c in model["columns"]], "rows": rows, "errors": errors}

    def start_export(self, pk, body):
        row = self.control.mapping(pk)
        require(row["status"] == "approved" and row["config_path"], "Approve and publish the mapping before exporting")
        config = load_config(row["config_path"])
        resume = bool(body.get("resume"))
        with self.lock:
            require(config["output_dir"] not in self.running, "An export for this mapping is already running")
            self.running.add(config["output_dir"])
        export_pk = self.control.start_export(pk, row["config_path"], config["output_dir"])

        def work():
            try:
                mapping = load_mapping(config["mapping"])
                with open_source(config, mapping) as source:
                    result = export(config, mapping, source, resume)
                self.control.finish_export(export_pk, result)
            except Exception as exc:
                self.control.finish_export(export_pk, error=str(exc) if isinstance(exc, ExportError) else "Export failed; run it from the CLI for details")
            finally:
                with self.lock:
                    self.running.discard(config["output_dir"])
        threading.Thread(target=work, daemon=True).start()
        return {"export_id": export_pk, "output_dir": config["output_dir"]}

    def exports(self, _):
        return self.control.exports()

    def start_job(self, body):
        provider = body.get("provider") or "devin"
        require(provider in {"devin", "codex"}, "Studio jobs use the devin or codex provider")
        document = None
        if body.get("document"):
            document = (self.root / "uploads" / Path(body["document"]).name).resolve()
            require(document.is_file() and document.parent == self.root / "uploads", "Unknown uploaded document")
        source_id, target_id = body.get("source_id"), body.get("target_id")
        metadata = self.control.source(source_id)["metadata"] if source_id else None
        spec = self.control.target(target_id)["spec"] if target_id else None
        started = start_job(document, self.root / "jobs", provider, timeout=int(body.get("timeout") or 1800),
                            source_metadata=metadata, target_spec=spec)
        pk = self.control.record_job(started["job_dir"], provider, "started", document.name if document else None,
                                     None, source_id, target_id)
        return {"id": pk} | started

    def job(self, pk):
        row = self.control.job(pk)
        status = status_job(row["job_dir"])
        self.control.record_job(row["job_dir"], row["provider"], status["status"], session_id=status.get("session_id"),
                                session_url=status.get("session_url"), result=status if status["status"] not in {"queued", "running"} else None)
        proposals = [load_mapping(o["path"])["model"] for o in status.get("outputs", [])] if status["status"] == "completed" else []
        return {"id": pk, "provider": row["provider"], "status": status["status"], "error": status.get("error"),
                "provider_status": status.get("provider_status"), "session_url": status.get("session_url"),
                "issues": status.get("issues", []), "proposals": proposals, "source_id": row["source_id"], "target_id": row["target_id"]}

    def cancel(self, pk, _):
        return cancel_job(self.control.job(pk)["job_dir"])

    def audit(self, _):
        return self.control.audit_log()


def handler(studio, token, port):
    allowed_hosts = {f"127.0.0.1:{port}", f"localhost:{port}"}
    routes = {
        ("GET", "state"): studio.state, ("POST", "connections"): studio.save_connection, ("POST", "tables"): studio.tables,
        ("POST", "describe"): studio.describe, ("POST", "sources"): studio.save_source, ("POST", "targets"): studio.save_target,
        ("POST", "automap"): studio.automap, ("POST", "validate"): studio.validate, ("POST", "mappings"): studio.save_mapping,
        ("POST", "preview"): studio.preview, ("GET", "exports"): studio.exports, ("POST", "jobs"): studio.start_job,
        ("GET", "audit"): studio.audit,
    }
    item_routes = {("GET", "mappings", None): lambda pk, _: studio.mapping(pk), ("POST", "mappings", "approve"): studio.approve,
                   ("POST", "mappings", "retire"): studio.retire, ("POST", "mappings", "dry-run"): studio.dry_run,
                   ("POST", "mappings", "export"): studio.start_export, ("GET", "jobs", None): lambda pk, _: studio.job(pk),
                   ("POST", "jobs", "cancel"): studio.cancel}

    class Handler(BaseHTTPRequestHandler):
        server_version = "MappingStudio"
        sys_version = ""

        def log_message(self, *args):
            pass  # Request lines can contain workspace names; keep the console quiet.

        def send(self, code, body, kind="application/json; charset=utf-8", headers=()):
            data = body if isinstance(body, bytes) else json.dumps(body, ensure_ascii=False).encode("utf-8")
            self.send_response(code)
            self.send_header("Content-Type", kind)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("X-Frame-Options", "DENY")
            self.send_header("Referrer-Policy", "no-referrer")
            self.send_header("Content-Security-Policy", "default-src 'self'; style-src 'self'; script-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'")
            for name, value in headers:
                self.send_header(name, value)
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self):
            self.dispatch("GET")

        def do_POST(self):
            self.dispatch("POST")

        def dispatch(self, method):
            if self.headers.get("Host") not in allowed_hosts:
                return self.send(403, {"error": "Studio only answers on 127.0.0.1 or localhost"})
            url = urlparse(self.path)
            if method == "GET" and url.path in ASSETS:
                name, kind = ASSETS[url.path]
                return self.send(200, (STATIC / name).read_bytes(), kind)
            parts = url.path.strip("/").split("/")
            if parts[0] != "api":
                return self.send(404, {"error": "Not found"})
            if not hmac.compare_digest(self.headers.get("X-Studio-Token", ""), token):
                return self.send(401, {"error": "Missing or wrong studio token; reopen the URL printed by 'studio serve'"})
            try:
                length = int(self.headers.get("Content-Length") or 0)
                require(0 <= length <= LIMIT, "Request body exceeds 10 MiB")
                raw = self.rfile.read(length) if length else b""
                query = {k: v[0] for k, v in parse_qs(url.query).items()}
                if parts[1:] == ["upload"] and method == "POST":
                    return self.send(200, studio.upload(query.get("kind"), query.get("name"), raw))
                if len(parts) == 4 and parts[1] == "mappings" and parts[3] == "download.xlsx" and method == "GET":
                    model = studio.mapping(int(parts[2]))["model"]
                    name = f"{model['metadata']['mapping_id']}-{model['metadata']['mapping_version']}.mapping.xlsx"
                    return self.send(200, workbook(model), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                     [("Content-Disposition", f'attachment; filename="{name}"')])
                body = json.loads(raw) if raw else {}
                require(isinstance(body, dict), "Request body must be a JSON object")
                if len(parts) == 2 and (method, parts[1]) in routes:
                    return self.send(200, routes[(method, parts[1])](body))
                if len(parts) in {3, 4} and parts[2].isdigit():
                    key = (method, parts[1], parts[3] if len(parts) == 4 else None)
                    if key in item_routes:
                        return self.send(200, item_routes[key](int(parts[2]), body))
                return self.send(404, {"error": "Unknown API route"})
            except ExportError as exc:
                return self.send(400, {"error": str(exc)})
            except (ValueError, KeyError, TypeError) as exc:
                return self.send(400, {"error": "Invalid request: " + type(exc).__name__})
            except Exception:
                return self.send(500, {"error": "Studio request failed; check the workspace and try again"})
    return Handler


def make_server(workspace, port=8765, token=None):
    studio = Studio(workspace)
    token = token or secrets.token_urlsafe(24)
    server = ThreadingHTTPServer(("127.0.0.1", port), handler(studio, token, port))
    if port == 0:
        # Re-bind the handler so the Host check uses the real ephemeral port.
        server.RequestHandlerClass = handler(studio, token, server.server_address[1])
    return server, token


def serve(workspace, port=8765, browser=True):
    server, token = make_server(workspace, port)
    # The token rides in the fragment so it never appears in request lines or proxy logs.
    url = f"http://127.0.0.1:{server.server_address[1]}/#token={token}"
    print(f"Mapping Studio for {Path(workspace).resolve()}\nOpen: {url}\nPress Ctrl+C to stop.", flush=True)
    if browser:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
